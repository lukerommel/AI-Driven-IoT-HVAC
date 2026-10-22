import time
from datetime import datetime
from pathlib import Path
from statistics import fmean
from typing import Optional, List
from collections import deque

import board
import busio
import adafruit_dht
import adafruit_ads1x15.ads1115 as ADS
from adafruit_ads1x15.analog_in import AnalogIn
from openpyxl import Workbook, load_workbook


# ------------------ USER SETTINGS ------------------
DHT_GPIO = board.D4
LOG_INTERVAL_SEC = 2.0

EXCEL_PATH = Path.home() / "sensor_readingsC.xlsx"
SHEET_NAME = "Readings"

# ADS1115 config
ADS_ADDR = 0x48
ADS_GAIN = 1
ADS_DATARATE = 860

# Channel mapping
VOLTAGE_CHANNEL = ADS.P0
CURRENT_CHANNEL = ADS.P1

# RMS sampling
SAMPLE_WINDOW_SEC = 0.20
OFFSET_SAMPLES = 40
TARGET_SAMPLE_HZ = 700

# ZMPT calibration
ZMPT_TO_MAINS_SCALE = 350.0
CALIBRATION_K = 230.0 / 115.0

# ACS712 calibration
ACS_SENSITIVITY = 0.030  # volts per amp

# Excel saving
SAVE_EVERY_ROWS = 10
SAVE_EVERY_SEC = 20

# Efficiency settings
SEI_SCALE_FACTOR = 1000
PERSISTENT_LIMIT = 3
LOW_LOAD_POWER_THRESHOLD = 100
HIGH_POWER_THRESHOLD = 500
# ---------------------------------------------------


HEADERS = [
    "Timestamp",
    "Temperature_C",
    "Humidity_%",
    "A0_Voltage_V",
    "Vrms_ADC",
    "Mains_Vrms_Est",
    "A1_Current_V",
    "Irms_A",
    "ApparentPower_VA",
    "Cooling_Rate_C_per_min",
    "Effective_Cooling_Rate_C_per_min",
    "Smart_Efficiency_Index",
    "CoolingRate_MA_5",
    "Power_MA_5",
    "SEI_MA_5",
    "Efficiency_Status",
    "Inefficiency_Cause",
    "Trend_Status",
    "Persistent_Inefficiency_Count",
    "Persistent_Warning",
    "Recommendation",
    "Predicted_Risk"
]


def setup_dht():
    return adafruit_dht.DHT22(DHT_GPIO, use_pulseio=False)


def setup_ads():
    i2c = busio.I2C(board.SCL, board.SDA)
    ads = ADS.ADS1115(i2c, address=ADS_ADDR)
    ads.gain = ADS_GAIN
    ads.data_rate = ADS_DATARATE
    return ads


def measure_rms(
    chan: AnalogIn,
    window_sec: float = SAMPLE_WINDOW_SEC,
    offset_samples: int = OFFSET_SAMPLES,
    target_hz: int = TARGET_SAMPLE_HZ
) -> Optional[float]:
    try:
        offs = [chan.voltage for _ in range(offset_samples)]
        v_offset = fmean(offs)

        dt = 1.0 / max(1, target_hz)
        t_end = time.perf_counter() + window_sec

        sum_sq = 0.0
        n = 0

        while True:
            if time.perf_counter() >= t_end:
                break
            v = chan.voltage - v_offset
            sum_sq += v * v
            n += 1
            time.sleep(dt)

        if n == 0:
            return None

        return (sum_sq / n) ** 0.5

    except Exception:
        return None


class ExcelLogger:
    def __init__(self, path: Path, sheet_name: str):
        self.path = path
        self.sheet_name = sheet_name
        self.buffer: List[List] = []
        self.last_save_ts = time.time()
        self._open_or_create()

    def _open_or_create(self):
        if self.path.exists():
            self.wb = load_workbook(self.path)
            if self.sheet_name in self.wb.sheetnames:
                self.ws = self.wb[self.sheet_name]
            else:
                self.ws = self.wb.create_sheet(self.sheet_name)

            if self.ws.max_row == 1 and self.ws.cell(1, 1).value is None:
                self.ws.append(HEADERS)
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = self.sheet_name
            self.ws.append(HEADERS)
            self.wb.save(self.path)

    def append(self, row: List):
        self.buffer.append(row)

    def flush_if_needed(self, force: bool = False):
        now = time.time()
        need_save = (
            force
            or len(self.buffer) >= SAVE_EVERY_ROWS
            or (now - self.last_save_ts) >= SAVE_EVERY_SEC
        )

        if not need_save or not self.buffer:
            return

        for r in self.buffer:
            self.ws.append(r)
        self.buffer.clear()

        for _ in range(3):
            try:
                self.wb.save(self.path)
                self.last_save_ts = time.time()
                return
            except PermissionError:
                time.sleep(0.5)

    def close(self):
        self.flush_if_needed(force=True)


def safe_mean(values):
    valid = [v for v in values if v is not None]
    if not valid:
        return None
    return sum(valid) / len(valid)


def calculate_live_cooling_rate(prev_temp, current_temp, dt_sec):
    if prev_temp is None or current_temp is None or dt_sec <= 0:
        return None
    dt_min = dt_sec / 60.0
    return (prev_temp - current_temp) / dt_min


def effective_cooling_rate(cooling_rate):
    if cooling_rate is None:
        return None
    return max(0, cooling_rate)


def calculate_sei(effective_cooling, apparent_power):
    if effective_cooling is None or apparent_power is None or apparent_power <= 0:
        return None
    return (effective_cooling / apparent_power) * SEI_SCALE_FACTOR


def classify_live_efficiency(power, cooling, sei, sei_avg):
    if power is None or sei is None:
        return "Unknown"

    if power < LOW_LOAD_POWER_THRESHOLD:
        return "Idle/Very Low Load"

    if cooling == 0 and power > HIGH_POWER_THRESHOLD:
        return "Inefficient"

    if sei_avg is not None and sei_avg > 0:
        if sei >= sei_avg * 1.05:
            return "Efficient"
        elif sei >= sei_avg * 0.85:
            return "Moderate"
        else:
            return "Inefficient"

    return "Moderate"


def diagnose_live_cause(power, power_avg, cooling, cooling_avg, sei, sei_avg):
    if power is None or sei is None:
        return "Insufficient data"

    if power < LOW_LOAD_POWER_THRESHOLD:
        return "Low-load or idle condition"

    if sei_avg is not None and sei_avg > 0 and sei >= sei_avg * 0.85:
        return "No major inefficiency detected"

    cooling_drop = (
        cooling_avg is not None
        and cooling_avg > 0
        and cooling is not None
        and cooling < cooling_avg * 0.85
    )
    power_rise = (
        power_avg is not None
        and power_avg > 0
        and power > power_avg * 1.10
    )

    if cooling_drop and power_rise:
        return "Cooling rate dropped and apparent power increased"
    elif cooling_drop:
        return "Cooling rate dropped"
    elif power_rise:
        return "Apparent power increased"
    else:
        return "Minor fluctuation or transitional behavior"


def get_live_trend(sei_history):
    valid = [v for v in sei_history if v is not None]
    if len(valid) < 3:
        return "Insufficient history"

    a, b, c = valid[-3], valid[-2], valid[-1]

    if c > b > a:
        return "Improving"
    elif c < b < a:
        return "Worsening"
    else:
        return "Stable/Fluctuating"


def get_live_predicted_risk(status, trend):
    if status == "Inefficient" and trend == "Worsening":
        return "High"
    elif status == "Inefficient":
        return "Medium"
    elif status == "Moderate":
        return "Low"
    else:
        return "Minimal"


def get_live_recommendation(status, cause, trend, persistent_warning):
    if status == "Efficient":
        return "System is operating efficiently."
    if status == "Moderate":
        return "Performance is acceptable but should be monitored."
    if status == "Idle/Very Low Load":
        return "System appears to be under very low load or idle condition."
    if status == "Unknown":
        return "Insufficient data for evaluation."

    if persistent_warning == "Yes":
        return "Persistent inefficiency detected across multiple cycles. System performance should be inspected."

    if cause == "Cooling rate dropped and apparent power increased":
        if trend == "Worsening":
            return "Combined degradation detected. Efficiency may fall further if this continues."
        return "Cooling is weaker while power draw is increasing."

    if cause == "Cooling rate dropped":
        return "Cooling effectiveness is reduced."

    if cause == "Apparent power increased":
        return "Power consumption increased without matching cooling improvement."

    if cause == "Low-load or idle condition":
        return "Low efficiency here may not indicate a fault."

    return "Monitor the system for further changes."


def main():
    print(f"Logging to: {EXCEL_PATH}")
    logger = ExcelLogger(EXCEL_PATH, SHEET_NAME)

    dht = setup_dht()
    ads = setup_ads()

    voltage_chan = AnalogIn(ads, VOLTAGE_CHANNEL)
    current_chan = AnalogIn(ads, CURRENT_CHANNEL)

    prev_temp = None
    prev_time = None

    cooling_history = deque(maxlen=5)
    power_history = deque(maxlen=5)
    sei_history = deque(maxlen=5)

    persistent_inefficiency_count = 0

    try:
        while True:
            now = datetime.now()

            # DHT22 reading
            try:
                t_c = dht.temperature
                rh = dht.humidity
            except RuntimeError:
                t_c = None
                rh = None

            # Voltage measurement
            v_instant = voltage_chan.voltage
            vrms_adc = measure_rms(voltage_chan)
            mains_vrms = (
                vrms_adc * ZMPT_TO_MAINS_SCALE * CALIBRATION_K
                if vrms_adc is not None else None
            )

            # Current measurement
            i_instant = current_chan.voltage
            irms_adc = measure_rms(current_chan)
            irms = (
                irms_adc / ACS_SENSITIVITY
                if irms_adc is not None else None
            )

            # Power
            apparent_power = (
                mains_vrms * irms
                if mains_vrms is not None and irms is not None else None
            )

            # Cooling rate
            if prev_time is not None and prev_temp is not None and t_c is not None:
                dt_sec = (now - prev_time).total_seconds()
                cooling_rate = calculate_live_cooling_rate(prev_temp, t_c, dt_sec)
            else:
                cooling_rate = None

            eff_cooling = effective_cooling_rate(cooling_rate)
            sei = calculate_sei(eff_cooling, apparent_power)

            # Update rolling histories
            cooling_history.append(eff_cooling)
            power_history.append(apparent_power)
            sei_history.append(sei)

            cooling_avg = safe_mean(cooling_history)
            power_avg = safe_mean(power_history)
            sei_avg = safe_mean(sei_history)

            # Live classification
            status = classify_live_efficiency(apparent_power, eff_cooling, sei, sei_avg)
            cause = diagnose_live_cause(apparent_power, power_avg, eff_cooling, cooling_avg, sei, sei_avg)
            trend = get_live_trend(sei_history)
            risk = get_live_predicted_risk(status, trend)

            if status == "Inefficient":
                persistent_inefficiency_count += 1
            else:
                persistent_inefficiency_count = 0

            persistent_warning = "Yes" if persistent_inefficiency_count >= PERSISTENT_LIMIT else "No"

            recommendation = get_live_recommendation(
                status, cause, trend, persistent_warning
            )

            # Terminal display
            print("\n==== LIVE HVAC EFFICIENCY STATUS ====")
            print(f"Time: {now:%Y-%m-%d %H:%M:%S}")
            print(f"Temperature: {t_c:.2f} C" if t_c is not None else "Temperature: --")
            print(f"Humidity: {rh:.2f} %" if rh is not None else "Humidity: --")
            print(f"Mains Vrms: {mains_vrms:.2f} V" if mains_vrms is not None else "Mains Vrms: --")
            print(f"Irms: {irms:.3f} A" if irms is not None else "Irms: --")
            print(f"Apparent Power: {apparent_power:.2f} VA" if apparent_power is not None else "Apparent Power: --")
            print(f"Cooling Rate: {cooling_rate:.3f} C/min" if cooling_rate is not None else "Cooling Rate: --")
            print(f"Effective Cooling Rate: {eff_cooling:.3f} C/min" if eff_cooling is not None else "Effective Cooling Rate: --")
            print(f"Smart Efficiency Index: {sei:.3f}" if sei is not None else "Smart Efficiency Index: --")
            print(f"Efficiency Status: {status}")
            print(f"Cause: {cause}")
            print(f"Trend: {trend}")
            print(f"Persistent Warning: {persistent_warning}")
            print(f"Recommendation: {recommendation}")
            print(f"Predicted Risk: {risk}")

            # Save row
            logger.append([
                now.strftime("%Y-%m-%d %H:%M:%S"),
                round(t_c, 2) if t_c is not None else None,
                round(rh, 2) if rh is not None else None,
                round(v_instant, 4) if v_instant is not None else None,
                round(vrms_adc, 4) if vrms_adc is not None else None,
                round(mains_vrms, 2) if mains_vrms is not None else None,
                round(i_instant, 4) if i_instant is not None else None,
                round(irms, 3) if irms is not None else None,
                round(apparent_power, 2) if apparent_power is not None else None,
                round(cooling_rate, 3) if cooling_rate is not None else None,
                round(eff_cooling, 3) if eff_cooling is not None else None,
                round(sei, 3) if sei is not None else None,
                round(cooling_avg, 3) if cooling_avg is not None else None,
                round(power_avg, 3) if power_avg is not None else None,
                round(sei_avg, 3) if sei_avg is not None else None,
                status,
                cause,
                trend,
                persistent_inefficiency_count,
                persistent_warning,
                recommendation,
                risk
            ])

            logger.flush_if_needed()

            prev_temp = t_c
            prev_time = now

            time.sleep(LOG_INTERVAL_SEC)

    except KeyboardInterrupt:
        print("\nStopped by user. Saving data...")
    finally:
        logger.close()
        try:
            dht.exit()
        except Exception:
            pass


if __name__ == "__main__":
    main()
    